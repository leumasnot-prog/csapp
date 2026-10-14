"""Importação da lista de classificação (CSV ou XLSX).

Colunas (cabeçalho na 1ª linha; acentos e maiúsculas são ignorados):

    posicao | nome | cpf | habilitacoes | email (opcional) | telefone (opcional)

`habilitacoes` traz itens separados por "|", cada item no formato
ETAPA:Disciplina. Exemplo: `EF1:Polivalente | EF2:Matemática`.
Etapas aceitas: EI, EF1, EF2, EJA, EE (e sinônimos, como "Anos iniciais").

O processo é em duas fases. Primeiro a validação (dry-run), que devolve erros
e avisos por linha. Depois a confirmação, que só grava se não houver nenhum
erro. É tudo ou nada.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field

from sqlalchemy import select, update
from sqlalchemy.orm import Session

from .. import audit
from ..models import (
    Classificacao,
    Disciplina,
    Etapa,
    Habilitacao,
    ListaClassificacao,
    Professor,
    Usuario,
)
from .cpf import chave, cpf_valido, somente_digitos

MAX_LINHAS = 20000

ALIASES_COLUNA = {
    "posicao": {"posicao", "classificacao", "pos", "colocacao", "ordem", "n", "no"},
    "nome": {"nome", "nome completo", "candidato", "professor"},
    "cpf": {"cpf"},
    "habilitacoes": {"habilitacoes", "habilitacao", "areas", "area"},
    "email": {"email", "e-mail", "e mail"},
    "telefone": {"telefone", "celular", "fone", "whatsapp"},
}
OBRIGATORIAS = ("posicao", "nome", "cpf", "habilitacoes")

ALIASES_ETAPA = {
    Etapa.EI: {"ei", "infantil", "educacao infantil", "creche", "pre-escola", "pre escola"},
    Etapa.EF1: {"ef1", "ef i", "ef 1", "fundamental i", "fundamental 1", "anos iniciais",
                "fundamental anos iniciais", "peb i", "peb 1"},
    Etapa.EF2: {"ef2", "ef ii", "ef 2", "fundamental ii", "fundamental 2", "anos finais",
                "fundamental anos finais", "peb ii", "peb 2"},
    Etapa.EJA: {"eja", "educacao de jovens e adultos"},
    Etapa.EE: {"ee", "educacao especial", "aee", "especial"},
}
_ETAPA_POR_CHAVE = {alias: etapa for etapa, aliases in ALIASES_ETAPA.items() for alias in aliases}


@dataclass
class Problema:
    linha: int
    campo: str
    mensagem: str


@dataclass
class LinhaValida:
    linha: int
    posicao: int
    nome: str
    cpf: str
    email: str | None
    telefone: str | None
    habilitacoes: list[tuple[Etapa, str]]  # (etapa, nome da disciplina como no arquivo)


@dataclass
class ResultadoImportacao:
    ok: bool = False
    total_linhas: int = 0
    erros: list[Problema] = field(default_factory=list)
    avisos: list[Problema] = field(default_factory=list)
    professores_novos: int = 0
    professores_atualizados: int = 0
    disciplinas_novas: list[str] = field(default_factory=list)
    lista_id: int | None = None
    previa: list[dict] = field(default_factory=list)


class ArquivoInvalido(Exception):
    pass


# ------------------------------------------------------------------ leitura

def _ler_csv(conteudo: bytes) -> list[list[str]]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = conteudo.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    amostra = texto[:4096]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t")
        delim = dialeto.delimiter
    except csv.Error:
        delim = ";" if amostra.count(";") >= amostra.count(",") else ","
    return [linha for linha in csv.reader(io.StringIO(texto), delimiter=delim)]


def _ler_xlsx(conteudo: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as e:
        raise ArquivoInvalido(f"Não foi possível ler a planilha: {e}") from e
    ws = wb.worksheets[0]
    linhas = []
    for row in ws.iter_rows(values_only=True):
        linhas.append(["" if v is None else _celula(v) for v in row])
    wb.close()
    return linhas


def _celula(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def ler_arquivo(nome: str, conteudo: bytes) -> list[list[str]]:
    nome = (nome or "").lower()
    if nome.endswith((".xlsx", ".xlsm")):
        return _ler_xlsx(conteudo)
    if nome.endswith((".csv", ".txt")):
        return _ler_csv(conteudo)
    raise ArquivoInvalido("Formato não suportado. Envie .csv ou .xlsx.")


# ------------------------------------------------------------------ validação

def _mapear_cabecalho(cabecalho: list[str]) -> dict[str, int]:
    mapa: dict[str, int] = {}
    for i, col in enumerate(cabecalho):
        k = chave(col).replace("º", "").strip()
        for campo, aliases in ALIASES_COLUNA.items():
            if k in aliases and campo not in mapa:
                mapa[campo] = i
    return mapa


def _parse_habilitacoes(txt: str) -> tuple[list[tuple[Etapa, str]], list[str]]:
    itens, erros = [], []
    for bruto in re.split(r"[|\n]", txt or ""):
        bruto = bruto.strip()
        if not bruto:
            continue
        m = re.match(r"^\s*([^:]+?)\s*[:\-–]\s*(.+?)\s*$", bruto)
        if not m:
            erros.append(f"'{bruto}' não está no formato ETAPA:Disciplina")
            continue
        etapa = _ETAPA_POR_CHAVE.get(chave(m[1]))
        if etapa is None:
            erros.append(f"etapa desconhecida '{m[1]}' (use EI, EF1, EF2, EJA ou EE)")
            continue
        itens.append((etapa, re.sub(r"\s+", " ", m[2]).strip()))
    return itens, erros


def validar(linhas: list[list[str]]) -> tuple[list[LinhaValida], ResultadoImportacao]:
    res = ResultadoImportacao()
    linhas = [l for l in linhas if any((c or "").strip() for c in l)]
    if not linhas:
        res.erros.append(Problema(0, "arquivo", "Arquivo vazio."))
        return [], res
    mapa = _mapear_cabecalho(linhas[0])
    faltando = [c for c in OBRIGATORIAS if c not in mapa]
    if faltando:
        res.erros.append(Problema(1, "cabeçalho", "Colunas obrigatórias ausentes: "
                                  + ", ".join(faltando)))
        return [], res
    dados = linhas[1:]
    res.total_linhas = len(dados)
    if len(dados) > MAX_LINHAS:
        res.erros.append(Problema(0, "arquivo", f"Máximo de {MAX_LINHAS} linhas."))
        return [], res

    def campo(row, nome):
        i = mapa.get(nome)
        return (row[i] if i is not None and i < len(row) else "").strip()

    validas: list[LinhaValida] = []
    vistos_cpf: dict[str, int] = {}
    vistas_pos: dict[int, int] = {}
    for n, row in enumerate(dados, start=2):
        erros_linha = 0

        def erro(c, m):
            nonlocal erros_linha
            erros_linha += 1
            res.erros.append(Problema(n, c, m))

        pos_txt = campo(row, "posicao").rstrip("ºo°.")
        try:
            posicao = int(pos_txt)
            if posicao <= 0:
                raise ValueError
        except ValueError:
            erro("posicao", f"posição inválida: '{pos_txt}'")
            posicao = None
        if posicao is not None:
            if posicao in vistas_pos:
                erro("posicao", f"posição {posicao} repetida (linha {vistas_pos[posicao]})")
            vistas_pos.setdefault(posicao, n)

        nome = re.sub(r"\s+", " ", campo(row, "nome"))
        if len(nome) < 3:
            erro("nome", "nome ausente ou muito curto")

        cpf = somente_digitos(campo(row, "cpf"))
        if len(cpf) < 11 and cpf:
            cpf = cpf.zfill(11)  # Excel costuma comer zeros à esquerda
        if not cpf_valido(cpf):
            erro("cpf", f"CPF inválido: '{campo(row, 'cpf')}'")
        elif cpf in vistos_cpf:
            erro("cpf", f"CPF repetido (linha {vistos_cpf[cpf]})")
        else:
            vistos_cpf[cpf] = n

        habs, erros_hab = _parse_habilitacoes(campo(row, "habilitacoes"))
        for e in erros_hab:
            erro("habilitacoes", e)
        if not habs and not erros_hab:
            erro("habilitacoes", "nenhuma habilitação informada")

        email = campo(row, "email") or None
        if email and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
            res.avisos.append(Problema(n, "email", f"e-mail inválido ignorado: '{email}'"))
            email = None
        telefone = somente_digitos(campo(row, "telefone")) or None
        if telefone and len(telefone) not in (10, 11, 12, 13):
            res.avisos.append(Problema(n, "telefone", f"telefone com formato incomum: '{telefone}'"))

        if erros_linha == 0:
            validas.append(LinhaValida(n, posicao, nome, cpf, email, telefone, habs))

    if vistas_pos and not res.erros:
        esperado = set(range(1, max(vistas_pos) + 1))
        lacunas = sorted(esperado - set(vistas_pos))
        if lacunas:
            amostra = ", ".join(map(str, lacunas[:10])) + ("…" if len(lacunas) > 10 else "")
            res.avisos.append(Problema(0, "posicao", f"posições ausentes na sequência: {amostra}"))
    validas.sort(key=lambda l: l.posicao)
    return validas, res


# ------------------------------------------------------------------ importação

def importar(s: Session, *, usuario: Usuario, nome_arquivo: str, conteudo: bytes,
             nome_lista: str, edital: str | None, ativar: bool, confirmar: bool,
             ip: str | None = None) -> ResultadoImportacao:
    linhas = ler_arquivo(nome_arquivo, conteudo)
    validas, res = validar(linhas)

    disciplinas = {d.chave: d for d in s.execute(select(Disciplina)).scalars()}
    novas: dict[str, str] = {}
    for l in validas:
        for _, nome_disc in l.habilitacoes:
            k = chave(nome_disc)
            if k not in disciplinas and k not in novas:
                novas[k] = nome_disc
    res.disciplinas_novas = sorted(novas.values())
    for nome_disc in res.disciplinas_novas:
        res.avisos.append(Problema(0, "habilitacoes",
                                   f"disciplina nova será cadastrada: '{nome_disc}'"))

    existentes = {p.cpf: p for p in s.execute(
        select(Professor).where(Professor.cpf.in_([l.cpf for l in validas]))).scalars()}
    res.professores_novos = sum(1 for l in validas if l.cpf not in existentes)
    res.professores_atualizados = len(validas) - res.professores_novos
    res.previa = [
        {"posicao": l.posicao, "nome": l.nome, "cpf_final": l.cpf[-4:],
         "habilitacoes": [f"{e.value}:{d}" for e, d in l.habilitacoes],
         "novo": l.cpf not in existentes}
        for l in validas[:50]
    ]
    res.ok = not res.erros and bool(validas)
    if not res.ok or not confirmar:
        return res

    # ---- gravação (tudo na transação de quem chamou)
    for k, nome_disc in novas.items():
        d = Disciplina(nome=nome_disc, chave=k)
        s.add(d)
        disciplinas[k] = d
    s.flush()

    lista = ListaClassificacao(
        nome=nome_lista.strip(), edital=(edital or "").strip() or None, ativa=False,
        arquivo_nome=nome_arquivo[:255], arquivo_sha256=hashlib.sha256(conteudo).hexdigest(),
        importada_por_id=usuario.id,
    )
    s.add(lista)
    s.flush()

    for l in validas:
        prof = existentes.get(l.cpf)
        if prof is None:
            prof = Professor(nome=l.nome, cpf=l.cpf, email=l.email, telefone=l.telefone)
            s.add(prof)
        else:
            prof.nome = l.nome
            prof.email = l.email or prof.email
            prof.telefone = l.telefone or prof.telefone
        prof.habilitacoes.clear()
        s.flush()
        for etapa, disciplina_id in dict.fromkeys(
                (e, disciplinas[chave(d)].id) for e, d in l.habilitacoes):
            prof.habilitacoes.append(Habilitacao(etapa=etapa, disciplina_id=disciplina_id))
        s.add(Classificacao(lista_id=lista.id, professor=prof, posicao=l.posicao))
    s.flush()

    if ativar:
        ativar_lista(s, lista, usuario, ip=ip, auditar=False)
    res.lista_id = lista.id
    audit.registrar(s, "LISTA_IMPORTADA", "lista_classificacao", lista.id, ator_tipo="USUARIO",
                    ator_id=usuario.id, ip=ip, dados={
                        "nome": lista.nome, "edital": lista.edital, "arquivo": nome_arquivo,
                        "sha256": lista.arquivo_sha256, "linhas": len(validas),
                        "professores_novos": res.professores_novos,
                        "professores_atualizados": res.professores_atualizados,
                        "disciplinas_novas": res.disciplinas_novas, "ativada": ativar,
                    })
    return res


def ativar_lista(s: Session, lista: ListaClassificacao, usuario: Usuario,
                 ip: str | None = None, auditar: bool = True) -> None:
    s.execute(update(ListaClassificacao).where(ListaClassificacao.id != lista.id)
              .values(ativa=False))
    s.flush()
    lista.ativa = True
    s.flush()
    if auditar:
        audit.registrar(s, "LISTA_ATIVADA", "lista_classificacao", lista.id,
                        ator_tipo="USUARIO", ator_id=usuario.id, ip=ip)
